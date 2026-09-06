"""
Артефакты — «как у Claude», но локально и в песочнице.
Модель сохраняет готовые файлы (текст, markdown, презентации)
ТОЛЬКО в ~/Luna/artifacts — вне песочницы запись невозможна,
поэтому инструмент безопасен и не требует подтверждения.
"""
import os
import re
import logging
import html as _html
from datetime import datetime

log = logging.getLogger("secretary.artifacts")

ARTIFACTS_DIR = os.path.expanduser("~/Luna/artifacts")


def _ensure_dir():
    os.makedirs(ARTIFACTS_DIR, exist_ok=True)


def safe_name(title: str) -> str:
    name = re.sub(r'[^a-zA-Zа-яА-Я0-9 _-]', '', title).strip().replace(' ', '')
    return name or "artifact"


def _parse_slides(content: str):
    """'# ' = заголовок слайда, '- ' = пункт."""
    slides = []
    cur = None
    for line in content.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith('#'):
            cur = {"heading": s.lstrip('#').strip(), "points": []}
            slides.append(cur)
        else:
            if cur is None:
                cur = {"heading": " ", "points": []}
                slides.append(cur)
            cur["points"].append(s.lstrip('-*').strip())
    return slides


def save_artifact(title: str, content: str, kind: str = "md") -> str:
    """Сохраняет готовый файл в песочницу артефактов.

    Args:
        title: название файла (без расширения)
        content: содержимое файла
        kind: md | txt | html
    """
    try:
        if kind not in ("md", "txt", "html"):
            kind = "md"
        _ensure_dir()
        stamp = datetime.now().strftime("%Y%m%d_%H%M")
        path = os.path.join(ARTIFACTS_DIR, f"{safe_name(title)}_{stamp}.{kind}")
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        log.info(f"Артефакт сохранён: {path}")
        return f"Сохранила файл: {path}"
    except Exception as e:
        log.error(f"Ошибка артефакта: {e}")
        return f"Не смогла сохранить файл: {e}"


def create_presentation(title: str, content: str) -> str:
    """Собирает презентацию из простого markdown и кладёт в песочницу.
    Формат content: '# Заголовок слайда', затем '- пункт' для каждого слайда.
    Всегда создаёт .html (открывается в браузере, листается стрелками/кликом)
    и .pptx тоже, если установлен python-pptx (pip install python-pptx).

    Args:
        title: название презентации
        content: слайды в простом markdown
    """
    try:
        _ensure_dir()
        stamp = datetime.now().strftime("%Y%m%d_%H%M")
        base = os.path.join(ARTIFACTS_DIR, f"{safe_name(title)}_{stamp}")
        slides = _parse_slides(content) or [{"heading": title, "points": []}]
        with open(base + ".html", "w", encoding="utf-8") as f:
            f.write(_render_html(title, slides))
        made = [base + ".html"]
        try:
            from pptx import Presentation
            prs = Presentation()
            for sl in slides:
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = sl["heading"]
                body = slide.placeholders[1].text_frame
                body.clear()
                for i, p in enumerate(sl["points"]):
                    tf = body if i == 0 else body.add_paragraph()
                    tf.text = p
            prs.save(base + ".pptx")
            made.append(base + ".pptx")
        except ImportError:
            log.info("python-pptx не установлен — только HTML-версия")
        return "Готово! Презентация: " + ", ".join(made)
    except Exception as e:
        log.error(f"Ошибка презентации: {e}")
        return f"Не смогла собрать презентацию: {e}"


def _render_html(title, slides):
    parts = []
    for sl in slides:
        pts = " ".join(f"<li>{_html.escape(p)}</li>" for p in sl["points"])
        parts.append(
            f'<section class="slide" style="display:none">'
            f'<h2>{_html.escape(sl["heading"])}</h2><ul>{pts}</ul></section>'
        )
    body = "\n".join(parts)
    return f"""<!DOCTYPE html>
<html lang="ru"><head><meta charset="utf-8"><title>{_html.escape(title)}</title>
<style>
body{{font-family:sans-serif;background:#101418;color:#eee;margin:0}}
.slide{{min-height:100vh;display:flex;flex-direction:column;justify-content:center;padding:8vw}}
h2{{color:#8AB4F8;font-size:6vw}}
li{{font-size:3.4vw;margin:.4em 0}}
#nav{{position:fixed;bottom:12px;right:16px;color:#667}}
</style></head><body>
{body}
<div id="nav"></div>
<script>
const s=[...document.querySelectorAll('.slide')];let i=0;
function show(){{s.forEach((e,j)=>e.style.display=j===i?'flex':'none');document.getElementById('nav').textContent=(i+1)+' / '+s.length;}}
document.addEventListener('keydown',e=>{{if(e.key==='ArrowRight')i=Math.min(i+1,s.length-1);if(e.key==='ArrowLeft')i=Math.max(i-1,0);show();}});
document.addEventListener('click',()=>{{i=Math.min(i+1,s.length-1);show();}});
show();
</script></body></html>"""
